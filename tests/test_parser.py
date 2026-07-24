import io

import pytest
from docx import Document

from app.services.chunker import chunk_document
from app.services.parser import ParseError, parse_file_bytes


def test_markdown_keeps_heading_locators():
    parsed = parse_file_bytes(
        "# 결정\nA안을 확정했다.\n\n## 승인\n김 팀장이 승인한다.".encode(),
        "decision.md",
    )
    assert parsed.mime_type == "text/markdown"
    assert [section.locator for section in parsed.sections] == ["결정", "승인"]
    assert "A안을 확정했다" in parsed.text


def test_docx_parses_paragraphs_and_table():
    stream = io.BytesIO()
    document = Document()
    document.add_heading("회의 결정", level=1)
    document.add_paragraph("B안 대신 A안을 사용한다.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "승인자"
    table.cell(0, 1).text = "김 팀장"
    document.save(stream)

    parsed = parse_file_bytes(stream.getvalue(), "회의록.docx")
    assert "A안을 사용한다" in parsed.text
    assert "김 팀장" in parsed.text
    assert any(section.locator == "표 1" for section in parsed.sections)


def test_rejects_disguised_pdf():
    with pytest.raises(ParseError, match="PDF 서명"):
        parse_file_bytes(b"not a pdf", "fake.pdf")


def test_chunker_splits_large_sections_with_overlap():
    parsed = parse_file_bytes(("결정 내용입니다. " * 300).encode(), "long.txt")
    chunks = chunk_document(parsed, target_chars=300, overlap_chars=40)
    assert len(chunks) > 2
    assert all(chunk.token_estimate > 0 for chunk in chunks)

