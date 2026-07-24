from __future__ import annotations

import io
import mimetypes
import re
from dataclasses import dataclass, field
from pathlib import Path

from charset_normalizer import from_bytes
from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MIME_BY_EXTENSION = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}


class ParseError(ValueError):
    pass


@dataclass
class ParsedSection:
    text: str
    locator: str | None = None


@dataclass
class ParsedDocument:
    title: str
    text: str
    mime_type: str
    sections: list[ParsedSection]
    warnings: list[str] = field(default_factory=list)


def normalized_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def validate_supported(filename: str) -> str:
    extension = normalized_extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise ParseError(
            f"지원하지 않는 형식입니다: {extension or '확장자 없음'}. "
            "PDF, DOCX, TXT, MD만 업로드할 수 있습니다."
        )
    return extension


def detect_mime(filename: str, declared_mime: str | None = None) -> str:
    extension = validate_supported(filename)
    return MIME_BY_EXTENSION.get(extension) or declared_mime or mimetypes.guess_type(filename)[0] or "application/octet-stream"


def parse_file_bytes(data: bytes, filename: str, declared_mime: str | None = None) -> ParsedDocument:
    extension = validate_supported(filename)
    if not data:
        raise ParseError("빈 파일은 처리할 수 없습니다.")

    if extension == ".pdf":
        return _parse_pdf(data, filename)
    if extension == ".docx":
        return _parse_docx(data, filename)
    return _parse_text(data, filename, detect_mime(filename, declared_mime))


def _parse_pdf(data: bytes, filename: str) -> ParsedDocument:
    if not data.startswith(b"%PDF"):
        raise ParseError("확장자는 PDF이지만 실제 PDF 서명이 없습니다.")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ParseError("암호화된 PDF는 MVP에서 지원하지 않습니다.") from exc
        sections: list[ParsedSection] = []
        empty_pages = 0
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(ParsedSection(text=text, locator=f"{page_number}페이지"))
            else:
                empty_pages += 1
        warnings = []
        if empty_pages:
            warnings.append(
                f"텍스트를 추출하지 못한 페이지가 {empty_pages}개 있습니다. 스캔 PDF는 OCR 고도화 대상입니다."
            )
        full_text = "\n\n".join(section.text for section in sections)
        if not full_text:
            raise ParseError("PDF에서 텍스트를 추출하지 못했습니다. 스캔 PDF일 가능성이 있습니다.")
        return ParsedDocument(
            title=Path(filename).stem,
            text=full_text,
            mime_type="application/pdf",
            sections=sections,
            warnings=warnings,
        )
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"PDF 파싱에 실패했습니다: {exc}") from exc


def _parse_docx(data: bytes, filename: str) -> ParsedDocument:
    if not data.startswith(b"PK"):
        raise ParseError("확장자는 DOCX이지만 올바른 Office 문서가 아닙니다.")
    try:
        document = DocxDocument(io.BytesIO(data))
        sections: list[ParsedSection] = []
        current_heading = "문서 시작"
        paragraph_buffer: list[str] = []

        def flush() -> None:
            nonlocal paragraph_buffer
            text = "\n".join(paragraph_buffer).strip()
            if text:
                sections.append(ParsedSection(text=text, locator=current_heading))
            paragraph_buffer = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                flush()
                current_heading = text
            else:
                paragraph_buffer.append(text)
        flush()

        for table_index, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            table_text = "\n".join(rows).strip()
            if table_text:
                sections.append(ParsedSection(text=table_text, locator=f"표 {table_index}"))

        full_text = "\n\n".join(section.text for section in sections)
        if not full_text:
            raise ParseError("DOCX에서 읽을 수 있는 텍스트가 없습니다.")
        return ParsedDocument(
            title=Path(filename).stem,
            text=full_text,
            mime_type=MIME_BY_EXTENSION[".docx"],
            sections=sections,
        )
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"DOCX 파싱에 실패했습니다: {exc}") from exc


def _parse_text(data: bytes, filename: str, mime_type: str) -> ParsedDocument:
    best = from_bytes(data).best()
    if best is None:
        raise ParseError("문자 인코딩을 판별하지 못했습니다.")
    text = str(best).replace("\x00", "").strip()
    if not text:
        raise ParseError("문서에서 읽을 수 있는 텍스트가 없습니다.")

    sections: list[ParsedSection] = []
    if normalized_extension(filename) == ".md":
        heading = "문서 시작"
        buffer: list[str] = []
        for line in text.splitlines():
            if re.match(r"^#{1,6}\s+", line):
                section_text = "\n".join(buffer).strip()
                if section_text:
                    sections.append(ParsedSection(text=section_text, locator=heading))
                heading = re.sub(r"^#{1,6}\s+", "", line).strip()
                buffer = []
            else:
                buffer.append(line)
        section_text = "\n".join(buffer).strip()
        if section_text:
            sections.append(ParsedSection(text=section_text, locator=heading))
    else:
        sections = [ParsedSection(text=text, locator="본문")]

    return ParsedDocument(
        title=Path(filename).stem,
        text=text,
        mime_type=mime_type,
        sections=sections,
    )

